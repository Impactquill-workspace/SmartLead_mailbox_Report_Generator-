"""
Excel report export utilities using openpyxl.

Generates:
  Sheet 1: Weekly Mailbox Report  (per-mailbox detail)
  Sheet 2: Client Summary         (per-client totals)
  Sheet 3: Report Metadata        (generation metadata)
"""
from __future__ import annotations

import logging
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# ---- Colour palette ----
HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
ALT_ROW_FILL = PatternFill(start_color="EBF3FF", end_color="EBF3FF", fill_type="solid")
NORMAL_FONT = Font(name="Calibri", size=10)
TITLE_FONT = Font(name="Calibri", bold=True, size=12)


def _apply_header_row(ws: Any, headers: list[str], row: int = 1) -> None:
    """Write bold header row with background colour."""
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _auto_size_columns(ws: Any, min_width: int = 10, max_width: int = 50) -> None:
    """Auto-size columns to fit content."""
    for col_cells in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        adjusted_width = min(max(max_length + 2, min_width), max_width)
        ws.column_dimensions[col_letter].width = adjusted_width


def _set_alternating_rows(ws: Any, start_row: int = 2) -> None:
    """Apply light-blue fill to alternating data rows."""
    for row_idx, row in enumerate(ws.iter_rows(min_row=start_row), start=0):
        if row_idx % 2 == 1:
            for cell in row:
                cell.fill = ALT_ROW_FILL


def generate_excel_report(
    report_data: dict[str, Any],
    period_start: date,
    period_end: date,
    output_dir: str | Path = "exports",
) -> Path:
    """
    Generate the weekly mailbox report Excel file.

    Args:
        report_data: Aggregated report data in the format:
            {
                "clients": [
                    {
                        "client_id": int,
                        "client_name": str,
                        "mailboxes": [
                            {
                                "mailbox_id": int | None,
                                "email": str,
                                "sent": int,
                                "bounced": int,
                                "bounce_rate": float,
                            }
                        ]
                    }
                ]
            }
        period_start: Start date of the reporting period.
        period_end: End date of the reporting period.
        output_dir: Directory where the Excel file will be saved.

    Returns:
        Path to the generated Excel file.
    """
    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    filename = (
        f"weekly_mailbox_report_{period_start.isoformat()}"
        f"_to_{period_end.isoformat()}.xlsx"
    )
    file_path = output_path / filename

    wb = openpyxl.Workbook()

    # ---- Sheet 1: Weekly Mailbox Report ----
    ws1 = wb.active
    ws1.title = "Weekly Mailbox Report"

    headers_s1 = [
        "Date Range",
        "Client Name",
        "Mailbox Email",
        "Leads Contacted",
        "Total Sent",
        "Bounced",
        "Bounced Percentage",
        "Client ID",
        "Mailbox ID",
    ]
    _apply_header_row(ws1, headers_s1, row=1)
    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = ws1.dimensions

    date_range_str = f"{period_start.isoformat()} to {period_end.isoformat()}"

    row_num = 2
    for client in report_data.get("clients", []):
        for mailbox in client.get("mailboxes", []):
            bounce_rate = mailbox.get("bounce_rate", 0.0)
            leads_cnt = mailbox.get("leads_contacted", 0)
            ws1.cell(row=row_num, column=1, value=date_range_str)
            ws1.cell(row=row_num, column=2, value=client.get("client_name", ""))
            ws1.cell(row=row_num, column=3, value=mailbox.get("email", ""))
            ws1.cell(row=row_num, column=4, value=leads_cnt)
            ws1.cell(row=row_num, column=5, value=mailbox.get("sent", 0))
            ws1.cell(row=row_num, column=6, value=mailbox.get("bounced", 0))

            # Bounced Percentage formatted e.g. 2.50%
            ws1.cell(row=row_num, column=7, value=f"{bounce_rate:.2f}%")

            ws1.cell(row=row_num, column=8, value=client.get("client_id"))
            ws1.cell(row=row_num, column=9, value=mailbox.get("mailbox_id"))

            for col in range(1, 10):
                ws1.cell(row=row_num, column=col).font = NORMAL_FONT
                ws1.cell(row=row_num, column=col).alignment = Alignment(
                    vertical="center"
                )

            row_num += 1

    _set_alternating_rows(ws1, start_row=2)
    _auto_size_columns(ws1)

    # ---- Sheet 2: Client Summary ----
    ws2 = wb.create_sheet("Client Summary")
    headers_s2 = [
        "Client Name",
        "Client ID",
        "Total Mailboxes",
        "Leads Contacted",
        "Total Sent",
        "Total Bounced",
        "Overall Bounce Rate (%)",
    ]
    _apply_header_row(ws2, headers_s2, row=1)
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = "A1:G1"

    row_num = 2
    for client in report_data.get("clients", []):
        mailboxes = client.get("mailboxes", [])
        total_leads = sum(m.get("leads_contacted", 0) for m in mailboxes)
        total_sent = sum(m.get("sent", 0) for m in mailboxes)
        total_bounced = sum(m.get("bounced", 0) for m in mailboxes)
        overall_rate = round((total_bounced / total_sent * 100) if total_sent > 0 else 0, 2)

        ws2.cell(row=row_num, column=1, value=client.get("client_name", ""))
        ws2.cell(row=row_num, column=2, value=client.get("client_id"))
        ws2.cell(row=row_num, column=3, value=len(mailboxes))
        ws2.cell(row=row_num, column=4, value=total_leads)
        ws2.cell(row=row_num, column=5, value=total_sent)
        ws2.cell(row=row_num, column=6, value=total_bounced)
        ws2.cell(row=row_num, column=7, value=overall_rate).number_format = "0.00"

        for col in range(1, 8):
            ws2.cell(row=row_num, column=col).font = NORMAL_FONT

        row_num += 1

    _set_alternating_rows(ws2, start_row=2)
    _auto_size_columns(ws2)

    # ---- Sheet 3: Report Metadata ----
    ws3 = wb.create_sheet("Report Metadata")
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 40

    title_cell = ws3.cell(row=1, column=1, value="Report Metadata")
    title_cell.font = TITLE_FONT

    metadata_rows = [
        ("Generated At", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Period Start", period_start.isoformat()),
        ("Period End", period_end.isoformat()),
        ("API Source", "Smartlead API v1 (https://server.smartlead.ai/api/v1)"),
        ("Number of Clients", len(report_data.get("clients", []))),
        (
            "Number of Mailboxes",
            sum(
                len(c.get("mailboxes", []))
                for c in report_data.get("clients", [])
            ),
        ),
        (
            "Total Emails Sent",
            sum(
                sum(m.get("sent", 0) for m in c.get("mailboxes", []))
                for c in report_data.get("clients", [])
            ),
        ),
        (
            "Total Emails Bounced",
            sum(
                sum(m.get("bounced", 0) for m in c.get("mailboxes", []))
                for c in report_data.get("clients", [])
            ),
        ),
        ("Data Attribution", report_data.get("data_attribution_note", "")),
    ]

    for idx, (key, value) in enumerate(metadata_rows, start=2):
        ws3.cell(row=idx, column=1, value=key).font = Font(bold=True, name="Calibri")
        ws3.cell(row=idx, column=2, value=str(value)).font = NORMAL_FONT

    try:
        wb.save(file_path)
    except PermissionError:
        timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        alt_filename = (
            f"weekly_mailbox_report_{period_start.isoformat()}"
            f"_to_{period_end.isoformat()}_{timestamp_str}.xlsx"
        )
        file_path = output_path / alt_filename
        wb.save(file_path)
        logger.warning(
            "Primary Excel file was locked/open. Saved with timestamp: %s",
            file_path,
        )

    logger.info("Excel report saved: %s", file_path)
    return file_path
